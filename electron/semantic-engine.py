"""Pigeon local multimodal indexing engine. JSON-lines protocol; stdout is protocol only."""
import os
for key in ("OMP_NUM_THREADS", "MKL_NUM_THREADS", "OPENBLAS_NUM_THREADS", "NUMEXPR_NUM_THREADS"):
    os.environ[key] = os.environ.get("PIGEON_SEMANTIC_THREADS", "1")
os.environ["TOKENIZERS_PARALLELISM"] = "false"
os.environ["CUDA_VISIBLE_DEVICES"] = ""
try:os.nice(15)
except (AttributeError,OSError):pass
import sys, json, sqlite3, hashlib, math, time, tempfile, subprocess, re, io, zipfile
from pathlib import Path
from contextlib import redirect_stdout
from concurrent.futures import ThreadPoolExecutor
from threading import Lock
PDF_LOCK = Lock()
MODEL_ID = "google/embeddinggemma-2"
MODEL_VERSION = "embeddinggemma2-768-f32-v1"
model = None
QUERY_CACHE = {}
DB = sqlite3.connect(sys.argv[1], timeout=10)
DB.execute("PRAGMA journal_mode=WAL")
DB.execute("PRAGMA synchronous=NORMAL")
DB.executescript("""
CREATE TABLE IF NOT EXISTS assets (id TEXT PRIMARY KEY, fingerprint TEXT NOT NULL, version TEXT NOT NULL, cursor INTEGER DEFAULT 0, complete INTEGER DEFAULT 0, error TEXT DEFAULT '', name TEXT DEFAULT '');
CREATE TABLE IF NOT EXISTS vectors (asset_id TEXT NOT NULL, part INTEGER NOT NULL, vector BLOB NOT NULL, detail TEXT NOT NULL, PRIMARY KEY(asset_id,part));
CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY,value TEXT NOT NULL);
""")
columns={row[1] for row in DB.execute('PRAGMA table_info(assets)')}
for name,definition in [('content_fingerprint',"TEXT DEFAULT ''"),('metadata',"TEXT DEFAULT '{}'"),('metadata_dirty','INTEGER DEFAULT 0')]:
    if name not in columns:DB.execute(f'ALTER TABLE assets ADD COLUMN {name} {definition}')
DB.execute("DELETE FROM vectors WHERE asset_id IN (SELECT id FROM assets WHERE version != ?)",(MODEL_VERSION,))
DB.execute("DELETE FROM assets WHERE version != ?",(MODEL_VERSION,)); DB.commit()
FFMPEG = os.environ.get('PIGEON_SEMANTIC_FFMPEG','ffmpeg')
TEXT_EXT = set('txt md markdown json jsonc yaml yml csv tsv xml html htm css js cjs mjs ts tsx jsx py sh bash lua conf ini toml log rst sql c cpp h hpp rs go java rb php swift kt tex srt vtt'.split())
def reply(value):
    sys.stdout.write(json.dumps(value,ensure_ascii=False,allow_nan=False,separators=(',',':'))+'\n');sys.stdout.flush()
def load_model():
    global model
    if model is not None:return model
    with redirect_stdout(sys.stderr):
        import torch
        from sentence_transformers import SentenceTransformer
        torch.set_num_threads(int(os.environ.get('PIGEON_SEMANTIC_THREADS','1')))
        torch.set_num_interop_threads(1)
        model=SentenceTransformer(MODEL_ID,device='cpu',local_files_only=True,revision='914f7f89142e33e77833254d9c9b90c3cef7303b',model_kwargs={'torch_dtype':torch.float32})
        model.max_seq_length=2048
    return model

def vector(value):
    with redirect_stdout(sys.stderr):
        m=load_model()
        v=m.encode(value,normalize_embeddings=True,show_progress_bar=False,processing_kwargs={'image':{'max_soft_tokens':140},'video':{'max_soft_tokens':140}},batch_size=1)
    import numpy as np
    v=np.asarray(v,dtype='<f4').reshape(-1)
    if len(v)!=768 or not np.isfinite(v).all():raise ValueError('Invalid model embedding')
    norm=float(np.linalg.norm(v))
    if norm<1e-8:raise ValueError('Empty model embedding')
    return v/norm

def media_duration(path):
    result=subprocess.run([FFMPEG,'-hide_banner','-threads','1','-i',path],capture_output=True,timeout=15)
    match=re.search(rb'Duration: (\d+):(\d+):(\d+(?:\.\d+)?)',result.stderr)
    if not match:raise ValueError('Could not determine media duration')
    h,m,s=map(float,match.groups());return h*3600+m*60+s

def media_input(path,kind,start=0):
    if kind=='video':
        from PIL import Image
        result=subprocess.run([FFMPEG,'-v','error','-threads','1','-ss',str(start),'-i',path,'-t','20','-vf','fps=0.2:start_time=0:round=up,scale=384:384:force_original_aspect_ratio=decrease','-threads','1','-filter_threads','1','-f','image2pipe','-vcodec','mjpeg','pipe:1'],capture_output=True,timeout=35)
        if result.returncode:raise ValueError(result.stderr.decode(errors='replace')[-500:])
        frames=[];data=result.stdout;pos=0
        while len(frames)<4:
            begin=data.find(b'\xff\xd8',pos);end=data.find(b'\xff\xd9',begin+2)
            if begin<0 or end<0:break
            frames.append(Image.open(io.BytesIO(data[begin:end+2])).convert('RGB'));pos=end+2
        if not frames:raise ValueError('No video frames could be decoded')
        return {'video':frames}
    import numpy as np
    result=subprocess.run([FFMPEG,'-v','error','-threads','1','-ss',str(start),'-i',path,'-t','20','-vn','-ac','1','-ar','16000','-threads','1','-f','f32le','pipe:1'],capture_output=True,timeout=35)
    if result.returncode:raise ValueError('Could not decode audio')
    return {'audio':np.frombuffer(result.stdout,dtype='<f4').copy()}

def text_contents(path,ext):
    if ext in TEXT_EXT:
        with open(path,'rb') as source:data=source.read(8*1024*1024+1)
        if len(data)>8*1024*1024:raise ValueError('Text exceeds 8 MB indexing limit')
        if b'\x00' in data[:8192]:raise ValueError('Binary file has no text content')
        return data.decode('utf-8',errors='replace')
    if ext=='docx':
        from docx import Document
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist())>64*1024*1024:raise ValueError('Document exceeds extraction limit')
        d=Document(path);return '\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
    if ext in ('xlsx','xlsm'):
        from openpyxl import load_workbook
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist())>64*1024*1024:raise ValueError('Spreadsheet exceeds extraction limit')
        book=load_workbook(path,read_only=True,data_only=True);lines=[];length=0
        try:
            for sheet in book:
                for row in sheet.iter_rows(values_only=True):
                    line=' '.join(str(v) for v in row if v is not None);lines.append(line);length+=len(line)
                    if length>8*1024*1024:raise ValueError('Spreadsheet exceeds extraction limit')
        finally:book.close()
        return '\n'.join(lines)
    return ''

def content_part(asset,part):
    path=asset['path'];kind=asset.get('kind');ext=Path(path).suffix.lower().lstrip('.');title=asset.get('name') or Path(path).name
    if part==0:
        metadata=' '.join([title,asset.get('note',''),' '.join(asset.get('tags',[])),asset.get('path','')])[:4000]
        return f'title: {title} | text: {metadata}',{'type':'metadata','label':'Filename, notes and tags'},False
    if kind=='image':
        if part>1:return None,None,True
        from PIL import Image,ImageOps
        image=Image.open(asset.get('imagePath') or path)
        image=ImageOps.exif_transpose(image).convert('RGB');image.thumbnail((768,768))
        return {'image':image},{'type':'image','label':'Image content'},True
    if kind in ('video','audio'):
        duration=media_duration(path);start=(part-1)*20
        if start>=duration:return None,None,True
        return media_input(path,kind,start),{'type':kind,'start':start,'end':min(start+20,duration),'label':f'{kind.title()} {start}s–{min(start+20,duration):.0f}s'},start+20>=duration
    if ext=='pdf':
        from pypdf import PdfReader
        if os.path.getsize(path)>128*1024*1024:raise ValueError('PDF exceeds extraction limit')
        reader=PdfReader(path)
        page=part-1
        if page>=len(reader.pages):return None,None,True
        text=(reader.pages[page].extract_text() or '')[:7000]
        if not text.strip():
            import pypdfium2 as pdfium
            pdf=pdfium.PdfDocument(path)
            try:
                pdfpage=pdf[page]
                try:
                    bitmap=pdfpage.render(scale=1)
                    try:image=bitmap.to_pil().copy();image.thumbnail((768,768))
                    finally:bitmap.close()
                finally:pdfpage.close()
            finally:pdf.close()
            return {'image':image},{'type':'pdf','page':page+1,'label':f'Scanned PDF page {page+1}'},page+1>=len(reader.pages)
        return f'title: {title} | text: {text}',{'type':'pdf','page':page+1,'label':f'PDF page {page+1}','snippet':text[:240]},page+1>=len(reader.pages)
    text=text_contents(path,ext)
    if not text:return None,{'type':'metadata-only','label':'No supported content extractor; metadata indexed'},True
    offset=(part-1)*1800
    if offset>=len(text):return None,None,True
    chunk=text[offset:offset+2200]
    return f'title: {title} | text: {chunk}',{'type':'text','offset':offset,'label':f'Text passage {part}','snippet':chunk[:240]},offset+2200>=len(text)

def vector_batch(values,memory_limit=0):
    import numpy as np
    with redirect_stdout(sys.stderr):
        m=load_model()
        try:
            with open('/proc/self/statm') as source:rss=int(source.read().split()[1])*os.sysconf('SC_PAGE_SIZE')
        except (OSError,IndexError,ValueError,AttributeError):rss=memory_limit
        # Visual/audio activations vary widely. Keep those inference microbatches at one.
        batch_size=2 if all(isinstance(v,str) for v in values) and memory_limit-rss>=1024**3 else 1
        matrix=m.encode(values,normalize_embeddings=True,show_progress_bar=False,processing_kwargs={'image':{'max_soft_tokens':140},'video':{'max_soft_tokens':140}},batch_size=batch_size)
    matrix=np.asarray(matrix,dtype='<f4').reshape(len(values),-1)
    if matrix.shape[1]!=768 or not np.isfinite(matrix).all():raise ValueError('Invalid model embeddings')
    norms=np.linalg.norm(matrix,axis=1)
    if np.any(norms<1e-8):raise ValueError('Empty model embedding')
    return matrix/norms[:,None]

def prepare_content(asset,part):
    try:
        stat=os.stat(asset['path'])
        if abs(stat.st_mtime*1000-float(asset.get('modified',stat.st_mtime*1000)))>2000 or stat.st_size!=int(asset.get('size',stat.st_size)):
            raise ValueError('Source changed; rescan the portfolio before indexing this file')
        if Path(asset['path']).suffix.lower()=='.pdf':
            # PDFium shares native global state even across separate documents.
            with PDF_LOCK:value,detail,done=content_part(asset,part)
        else:value,detail,done=content_part(asset,part)
        error=detail['label'] if detail and detail.get('type')=='metadata-only' else ''
        return value,detail,done,error
    except Exception as exc:
        return None,None,True,str(exc)[:500]

def sync_assets(records,removed_ids=()):
    for ident in removed_ids:
        DB.execute('DELETE FROM vectors WHERE asset_id=?',(ident,));DB.execute('DELETE FROM assets WHERE id=?',(ident,))
    for record in records:
        asset=record['asset'];ident=asset['id'];expected=record['fingerprint'];content=record.get('contentFingerprint','');metadata=json.dumps(asset)
        row=DB.execute('SELECT fingerprint,content_fingerprint,metadata FROM assets WHERE id=?',(ident,)).fetchone()
        same_content=bool(row and row[1] and row[1]==content);same_metadata=bool(row and row[0]==expected)
        if row and not row[1] and content:
            # Upgrade old fingerprints without discarding unchanged content vectors.
            try:
                old=json.loads(row[0]);new=json.loads(content)
                if len(old)==9 and len(new)==9 and not new[4]:
                    same_content=old[:4]==new[:4] and old[7:9]==new[7:9]
                    same_metadata=same_content and old[4:7]==[asset.get('name'),asset.get('note',''),asset.get('tags',[])]
            except (ValueError,TypeError):pass
        if row and same_content and same_metadata and row[2]==metadata:continue
        if not row or not same_content:
            DB.execute('DELETE FROM vectors WHERE asset_id=?',(ident,))
            DB.execute('INSERT OR REPLACE INTO assets(id,fingerprint,version,name,content_fingerprint,metadata) VALUES (?,?,?,?,?,?)',(ident,expected,MODEL_VERSION,asset.get('name',''),content,metadata))
        else:
            if not same_metadata:DB.execute('DELETE FROM vectors WHERE asset_id=? AND part=0',(ident,))
            DB.execute('UPDATE assets SET fingerprint=?,content_fingerprint=?,metadata=?,name=?,metadata_dirty=CASE WHEN ? THEN metadata_dirty ELSE 1 END WHERE id=?',(expected,content,metadata,asset.get('name',''),int(same_metadata),ident))
    DB.commit()

def prepare_index_job(job):
    asset=job['asset'];ident=asset['id'];expected=job['fingerprint']
    row=DB.execute('SELECT fingerprint,cursor,complete,error,metadata_dirty FROM assets WHERE id=?',(ident,)).fetchone()
    if not row or row[0]!=expected:
        DB.execute('DELETE FROM vectors WHERE asset_id=?',(ident,))
        DB.execute('INSERT OR REPLACE INTO assets(id,fingerprint,version,name) VALUES (?,?,?,?)',(ident,expected,MODEL_VERSION,asset.get('name','')));DB.commit();row=(expected,0,0,'',0)
    return asset,row

def index_batch(jobs,workers=2,steps=2,memory_limit=0):
    if len(jobs)>4:raise ValueError('Index batches are limited to four files')
    if len({job['asset']['id'] for job in jobs})!=len(jobs):raise ValueError('Duplicate asset in index batch')
    results={};workers=max(1,min(4,int(workers),len(jobs) or 1))
    # Worker threads only extract/decode. One model and the main-thread SQLite writer are shared.
    round_workers=workers
    for _ in range(max(1,min(2,int(steps)))):
        round_workers=workers
        if memory_limit:
            try:
                with open('/proc/self/statm') as source:rss=int(source.read().split()[1])*os.sysconf('SC_PAGE_SIZE')
                round_workers=min(workers,max(1,(memory_limit-rss-32*1024**2)//(96*1024**2)))
            except (OSError,IndexError,ValueError,AttributeError):round_workers=1
        # Re-evaluate after metadata inference loads the model and consumes its RAM.
        with ThreadPoolExecutor(max_workers=round_workers,thread_name_prefix='pigeon-extract') as pool:
            pending=[]
            for job in jobs:
                asset,row=prepare_index_job(job);ident=asset['id']
                if row[2] and not row[4]:results[ident]={'id':ident,'done':True,'cached':True,'error':row[3]};continue
                part=0 if row[4] else row[1]
                pending.append((asset,part,pool.submit(prepare_content,asset,part)))
            prepared=[(asset,part,*future.result()) for asset,part,future in pending]
        values=[item[2] for item in prepared if item[2] is not None]
        embeddings=iter(vector_batch(values,memory_limit)) if values else iter(())
        for asset,part,value,detail,done,error in prepared:
            ident=asset['id']
            if value is not None:
                v=next(embeddings);DB.execute('INSERT OR REPLACE INTO vectors VALUES (?,?,?,?)',(ident,part,v.tobytes(),json.dumps(detail)))
            saved=DB.execute('SELECT cursor,complete,error,metadata_dirty FROM assets WHERE id=?',(ident,)).fetchone()
            if part==0 and saved[3] and saved[0]>0:
                done=bool(saved[1]);error=error or saved[2]
                DB.execute('UPDATE assets SET metadata_dirty=0 WHERE id=?',(ident,))
            else:DB.execute('UPDATE assets SET cursor=?,complete=?,error=?,metadata_dirty=0 WHERE id=?',(part+1,int(done),error,ident))
            DB.commit()
            results[ident]={'id':ident,'done':done,'part':part,'error':error}
    return {'results':list(results.values()),'workers':round_workers,'info':handle({'action':'info'})}

def index_step(asset,expected):
    # Preserve the single-file protocol for older previews and persisted-index tests.
    return index_batch([{'asset':asset,'fingerprint':expected}],workers=1,steps=1)['results'][0]

def search(request):
    import numpy as np,heapq
    def progress(stage,completed=0,total=0):reply({'id':request.get('id'),'progress':{'stage':stage,'completed':completed,'total':total}})
    progress('Preparing query')
    allowed=request.get('assets',{})
    key=json.dumps([request.get('query'),request.get('assetId'),allowed.get(request.get('assetId')),request.get('sample')],sort_keys=True)
    if key in QUERY_CACHE:query=QUERY_CACHE[key]
    elif request.get('query'):
        query=vector('task: search result | query: '+request['query'][:4000])
    elif request.get('assetId') and DB.execute('SELECT 1 FROM vectors v JOIN assets a ON a.id=v.asset_id WHERE v.asset_id=? AND v.part>0 AND a.fingerprint=? AND a.complete=1',(request['assetId'],allowed.get(request['assetId']))).fetchone():
        rows=DB.execute('SELECT vector FROM vectors WHERE asset_id=? AND part>0 ',(request['assetId'],)).fetchall()
        query=np.stack([np.frombuffer(r[0],dtype='<f4') for r in rows])
    else:
        sample=request.get('sample')
        if not sample:raise ValueError('Enter a description or choose a sample file')
        vectors=[];part=1;done=False
        while not done:
            value,detail,done=content_part(sample,part)
            if value is not None:vectors.append(vector(value))
            part+=1
        if not vectors:
            value,detail,done=content_part(sample,0);vectors.append(vector(value))
        query=np.stack(vectors)
    QUERY_CACHE[key]=query
    while len(QUERY_CACHE)>8 or sum(v.nbytes for v in QUERY_CACHE.values())>64*1024*1024:QUERY_CACHE.pop(next(iter(QUERY_CACHE)))
    offset=max(0,int(request.get('offset',0)))
    minimum=max(-1,min(1,float(request.get('minimum',0.68))));limit=max(1,min(1000,int(request.get('limit',200))));best={}
    total=DB.execute('SELECT count(*) FROM vectors').fetchone()[0];completed=0
    progress('Comparing indexed content',0,total)
    cursor=DB.execute('SELECT v.asset_id,v.vector,v.detail,a.fingerprint FROM vectors v JOIN assets a ON a.id=v.asset_id')
    while True:
        rows=cursor.fetchmany(256)
        if not rows:break
        completed+=len(rows);progress('Comparing indexed content',completed,total)
        valid=[r for r in rows if r[0] in allowed and allowed[r[0]]==r[3] and r[0]!=request.get('assetId')]
        if not valid:continue
        matrix=np.stack([np.frombuffer(r[1],dtype='<f4') for r in valid]);scores=matrix@query if query.ndim==1 else np.max(matrix@query.T,axis=1)
        for r,score in zip(valid,scores):
            score=float(score)
            if score>=minimum and (r[0] not in best or score>best[r[0]]['score']):best[r[0]]={'id':r[0],'score':score,'match':json.loads(r[2])}
    return {'results':heapq.nlargest(offset+limit,best.values(),key=lambda r:r['score'])[offset:],'totalMatches':len(best),'indexed':DB.execute('SELECT count(*) FROM assets WHERE complete=1').fetchone()[0]}

def handle(req):
    action=req['action']
    if action=='info':
        return {'indexed':DB.execute('SELECT count(*) FROM assets WHERE complete=1 AND error="" AND metadata_dirty=0').fetchone()[0], 'limited':DB.execute('SELECT count(*) FROM assets WHERE complete=1 AND error!="" AND metadata_dirty=0').fetchone()[0], 'vectors':DB.execute('SELECT count(*) FROM vectors').fetchone()[0],'settings':{r[0]:json.loads(r[1]) for r in DB.execute('SELECT key,value FROM settings')},'model':MODEL_ID,'version':MODEL_VERSION}
    if action=='configure':
        for k,v in req.get('settings',{}).items():DB.execute('INSERT OR REPLACE INTO settings VALUES (?,?)',(k,json.dumps(v)))
        DB.commit();return handle({'action':'info'})
    if action=='sync-assets':
        sync_assets(req.get('records',[]),req.get('removedIds',[]));return handle({'action':'info'})
    if action=='prune':
        if req.get('records'):sync_assets(req['records'])
        allowed=req['assets']
        for ident,fp in DB.execute('SELECT id,fingerprint FROM assets').fetchall():
            if allowed.get(ident)!=fp:DB.execute('DELETE FROM vectors WHERE asset_id=?',(ident,));DB.execute('DELETE FROM assets WHERE id=?',(ident,))
        for ident,fp in allowed.items():
            DB.execute('INSERT OR IGNORE INTO assets(id,fingerprint,version) VALUES (?,?,?)',(ident,fp,MODEL_VERSION))
        DB.commit();return handle({'action':'info'})
    if action=='plan':return {'pending':[r[0] for r in DB.execute('SELECT id FROM assets WHERE complete=0 OR metadata_dirty=1')]}
    if action=='index':return index_step(req['asset'],req['fingerprint'])
    if action=='index_batch':return index_batch(req.get('jobs',[]),req.get('workers',2),req.get('steps',2),int(req.get('memoryLimit',0)))
    if action=='search':return search(req)
    if action=='check':
        v=vector('task: search result | query: a bird flying');return {'dimensions':len(v),'finite':True}
    raise ValueError('Unknown semantic action')
reply({'ready':True})
for line in sys.stdin:
    try:
        req=json.loads(line)
        result=handle(req);reply({'id':req['id'],'result':result})
    except Exception as exc:
        reply({'id':req.get('id') if isinstance(locals().get('req'),dict) else None,'error':str(exc)[:1000]})
